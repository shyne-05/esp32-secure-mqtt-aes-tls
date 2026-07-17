from docx import Document
from pathlib import Path
import re

src = Path(r'C:\doan2\report_github.docx')
out = Path(r'C:\doan2\bao_cao_do_an_github_cap_nhat.docx')
doc = Document(src)

def fix_text(text):
    # Replace machine-specific paths in both prose and code listings.
    text = text.replace('C:\\doan2', 'PROJECT_ROOT').replace('c:\\doan2', 'PROJECT_ROOT')
    text = text.replace('C:/doan2', 'PROJECT_ROOT').replace('c:/doan2', 'PROJECT_ROOT')
    return text

def fix_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = fix_text(run.text)

for p in doc.paragraphs:
    fix_paragraph(p)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                fix_paragraph(p)
for section in doc.sections:
    for p in section.header.paragraphs + section.footer.paragraphs:
        fix_paragraph(p)

# Add a public, reproducible source-of-code note near the end of the report.
doc.add_page_break()
doc.add_heading('LIÊN KẾT MÃ NGUỒN VÀ KHẢ NĂNG TÁI LẬP', level=1)
doc.add_paragraph(
    'Toàn bộ mã nguồn, cấu trúc bốn giai đoạn, script listener, script phân tích, '
    'cấu hình chứng thư mẫu và các tệp hỗ trợ của đồ án được nhóm công khai tại repository GitHub:'
)
p = doc.add_paragraph('https://github.com/shyne-05/esp32-secure-mqtt-aes-tls')
for run in p.runs:
    run.bold = True
doc.add_paragraph(
    'Repository là nguồn tham chiếu chính cho mã nguồn trong báo cáo. Khi tái lập, người đọc '
    'có thể tải repository, cài ESP-IDF và các dependency theo hướng dẫn, sau đó chọn từng thư mục '
    'stage1_baseline, stage2_aesgcm, stage3_tls hoặc stage4_aesgcm_tls để build và flash. '
    'Các đường dẫn trong phần cấu hình được trình bày theo dạng tương đối hoặc PROJECT_ROOT, '
    'do đó không phụ thuộc vào tên ổ đĩa hay thư mục trên máy của người đọc.'
)
doc.add_heading('Ghi chú về tệp bằng chứng', level=2)
doc.add_paragraph(
    'Các log, CSV, biểu đồ và hình xác thực trong báo cáo là kết quả thu thập từ lần chạy thực nghiệm '
    'của nhóm. Ảnh Telegram vẫn được chừa riêng để nhóm tự chèn ảnh chụp tài khoản thật khi cần; '
    'không sử dụng ảnh minh họa thay cho bằng chứng thông báo.'
)
doc.save(out)
print(out)
