from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

src = Path(r'C:\doan2\bao_cao_do_an_github_cap_nhat.docx')
out = Path(r'C:\doan2\bao_cao_do_an_github_cap_nhat_anh_day_du.docx')
assets = Path(r'C:\doan2\report_assets')
doc = Document(src)

items = [
    ('architecture.png', 'Kiến trúc tổng quát hệ thống ESP32-S3 – MQTT – listener.'),
    ('esp32_board.png', 'Bo mạch ESP32-S3-DevKitC-1.'),
    ('esp32_block.png', 'Sơ đồ khối ESP32-S3-DevKitC-1.'),
    ('esp32_pins.jpg', 'Sơ đồ chân ESP32-S3-DevKitC-1.'),
    ('dht11_sensor.jpg', 'Cảm biến DHT11 và các phần tử đo nhiệt độ, độ ẩm.'),
    ('dht11_breadboard.jpg', 'Ví dụ bố trí cảm biến DHT11 trên breadboard.'),
    ('dht11_wiring.jpg', 'Sơ đồ nối VCC, DATA, GND và điện trở kéo lên cho DHT11.'),
    ('security_layers.png', 'So sánh bảo vệ payload bằng AES-GCM và bảo vệ kênh bằng TLS.'),
    ('packet.png', 'Cấu trúc gói AES-GCM: nonce, ciphertext và authentication tag.'),
    ('threat_model.png', 'Mô hình mối đe dọa và các lớp kiểm soát.'),
    ('tls_handshake.png', 'Luồng bắt tay TLS và truyền MQTT qua kênh bảo mật.'),
    ('sequence.png', 'Trình tự xử lý từ cảm biến tới CSV.'),
    ('flash_stage1.png', 'Bằng chứng nạp firmware Stage 1.'),
    ('flash_stage2.png', 'Bằng chứng nạp firmware Stage 2.'),
    ('flash_stage3.png', 'Bằng chứng nạp firmware Stage 3.'),
    ('flash_stage4.png', 'Bằng chứng nạp firmware Stage 4.'),
    ('listener_stage1.png', 'Listener và dữ liệu Stage 1.'),
    ('listener_stage2.png', 'Listener giải mã AES-GCM Stage 2.'),
    ('listener_stage3.png', 'Listener nhận dữ liệu MQTTS Stage 3.'),
    ('listener_stage4.png', 'Listener nhận dữ liệu AES-GCM kết hợp TLS Stage 4.'),
    ('csv_stage1.png', 'Dữ liệu CSV Stage 1.'),
    ('csv_stage2.png', 'Dữ liệu CSV Stage 2.'),
    ('csv_stage3.png', 'Dữ liệu CSV Stage 3.'),
    ('csv_stage4.png', 'Dữ liệu CSV Stage 4.'),
    ('comparison_charts.png', 'Biểu đồ so sánh bốn giai đoạn.'),
    ('cert_info.png', 'Metadata chứng thư TLS; không hiển thị private key.'),
    ('tamper_test.png', 'Kiểm thử sửa ciphertext và từ chối do sai authentication tag.'),
]

doc.add_page_break()
doc.add_heading('PHỤ LỤC HÌNH ẢNH MINH HỌA VÀ BẰNG CHỨNG', level=1)
doc.add_paragraph(
    'Các hình dưới đây được chèn trực tiếp vào tài liệu Word để bảo đảm khi mở báo cáo độc lập '
    'vẫn hiển thị đầy đủ. Hình minh họa lý thuyết được ghi nguồn trong phần tài liệu tham khảo; '
    'hình log, CSV, flash và kiểm thử là bằng chứng từ quá trình thực nghiệm của nhóm.'
)

for filename, caption in items:
    path = assets / filename
    if not path.exists():
        continue
    try:
        doc.add_picture(str(path), width=Inches(5.9))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cp.runs:
            run.italic = True
            run.font.size = Pt(9)
    except Exception as exc:
        doc.add_paragraph(f'[Không thể chèn {filename}: {exc}]')

doc.add_page_break()
doc.add_heading('GHI CHÚ ẢNH TELEGRAM', level=1)
doc.add_paragraph(
    'Vị trí ảnh chụp Telegram thật vẫn được để riêng theo yêu cầu của nhóm. '
    'Nhóm chèn ảnh sau khi nhận thông báo trên tài khoản Telegram; không dùng ảnh minh họa hoặc ảnh dựng.'
)
doc.save(out)
print(out)
