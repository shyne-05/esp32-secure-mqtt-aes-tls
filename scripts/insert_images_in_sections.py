from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

src = Path(r'C:\doan2\bao_cao_do_an_github_cap_nhat.docx')
out = Path(r'C:\doan2\bao_cao_do_an_github_cap_nhat_anh_dung_vi_tri.docx')
assets = Path(r'C:\doan2\report_assets')
doc = Document(src)

def find_heading(starts):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            return p
    raise ValueError(f'Không tìm thấy mục: {starts}')

def insert_after(anchor, filename, caption, width=5.7):
    path = assets / filename
    if not path.exists():
        return anchor
    doc.add_picture(str(path), width=Inches(width))
    image_p = doc.paragraphs[-1]
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    anchor._p.addnext(image_p._p)
    anchor = image_p
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)
    anchor._p.addnext(cap._p)
    return cap

def insert_group(heading, images):
    anchor = find_heading(heading)
    for filename, caption, width in images:
        anchor = insert_after(anchor, filename, caption, width)

# Chèn ngay tại phần lý thuyết tương ứng.
insert_group('2.1.', [
    ('esp32_board.png', 'Hình 2.1. Bo mạch ESP32-S3-DevKitC-1 [TL-01].', 4.8),
    ('esp32_block.png', 'Hình 2.2. Sơ đồ khối ESP32-S3-DevKitC-1 [TL-01].', 5.8),
    ('esp32_pins.jpg', 'Hình 2.3. Sơ đồ chân ESP32-S3-DevKitC-1 và vị trí GPIO4 [TL-01].', 5.2),
])
insert_group('2.2.', [
    ('dht11_sensor.jpg', 'Hình 2.4. Cảm biến DHT11 và phần tử đo nhiệt độ, độ ẩm [TL-09].', 4.8),
    ('dht11_breadboard.jpg', 'Hình 2.5. Ví dụ bố trí DHT11 trên breadboard [TL-09].', 5.2),
    ('dht11_wiring.jpg', 'Hình 2.6. Sơ đồ nối VCC, DATA, GND và điện trở kéo lên [TL-09].', 5.4),
])
insert_group('2.3.', [
    ('architecture.png', 'Hình 2.7. Kiến trúc publish/subscribe của hệ thống ESP32-S3 – MQTT – listener.', 5.9),
])
insert_group('2.4.', [
    ('security_layers.png', 'Hình 2.8. So sánh bảo vệ payload bằng AES-GCM và bảo vệ kênh bằng TLS.', 5.9),
    ('packet.png', 'Hình 2.9. Cấu trúc gói AES-GCM: nonce || ciphertext || tag.', 5.7),
    ('tamper_test.png', 'Hình 2.10. Kiểm thử sửa ciphertext và từ chối do authentication tag không hợp lệ.', 5.7),
])
insert_group('2.5.', [
    ('tls_handshake.png', 'Hình 2.11. Luồng bắt tay TLS và truyền MQTT qua kênh bảo mật.', 5.9),
    ('cert_info.png', 'Hình 2.12. Metadata chứng thư TLS; private key không đưa vào báo cáo.', 5.5),
])

# Chèn ảnh đúng vào từng kịch bản thực nghiệm.
for heading, filename, caption in [
    ('5.2.', 'listener_stage1.png', 'Hình 5.1. Listener và dữ liệu Stage 1.'),
    ('5.3.', 'listener_stage2.png', 'Hình 5.2. Payload Base64 và giải mã AES-GCM Stage 2.'),
    ('5.4.', 'listener_stage3.png', 'Hình 5.3. Kết nối MQTTS và thời gian TLS Stage 3.'),
    ('5.5.', 'listener_stage4.png', 'Hình 5.4. Kết hợp AES-GCM và TLS ở Stage 4.'),
]:
    insert_group(heading, [(filename, caption, 5.9)])

insert_group('5.6.', [
    ('csv_stage1.png', 'Hình 5.5. CSV Stage 1 sau khi thu thập dữ liệu.', 5.8),
    ('csv_stage2.png', 'Hình 5.6. CSV Stage 2 sau khi listener xác thực và giải mã.', 5.8),
    ('csv_stage3.png', 'Hình 5.7. CSV Stage 3 sau khi nhận dữ liệu MQTTS.', 5.8),
    ('csv_stage4.png', 'Hình 5.8. CSV Stage 4 sau khi xử lý hai lớp bảo mật.', 5.8),
])

insert_group('6.2.', [
    ('comparison_charts.png', 'Hình 6.1. Biểu đồ so sánh kích thước payload, AES, TLS và heap.', 6.0),
])

# Các phiếu xác thực được đặt ảnh ngay dưới tiêu đề phiếu tương ứng.
for heading, filename, caption in [
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 1', 'flash_stage1.png', 'Hình XN-01. Ảnh xác thực flash Stage 1.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 2', 'csv_stage1.png', 'Hình XN-02. Ảnh xác thực CSV Stage 1.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 3', 'listener_stage2.png', 'Hình XN-03. Ảnh xác thực listener Stage 2.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 4', 'csv_stage2.png', 'Hình XN-04. Ảnh xác thực CSV Stage 2.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 5', 'listener_stage3.png', 'Hình XN-05. Ảnh xác thực listener Stage 3.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 6', 'listener_stage4.png', 'Hình XN-06. Ảnh xác thực listener Stage 4.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 7', 'comparison_charts.png', 'Hình XN-07. Ảnh xác thực biểu đồ tổng hợp.'),
    ('PHIẾU XÁC THỰC THỰC NGHIỆM 8', 'cert_info.png', 'Hình XN-08. Ảnh xác thực metadata chứng thư TLS.'),
]:
    insert_group(heading, [(filename, caption, 5.9)])

doc.save(out)
print(out)
